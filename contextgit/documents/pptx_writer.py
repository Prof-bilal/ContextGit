"""Render DocumentContent to a designed 16:9 `.pptx` deck.

Slides are drawn on a blank layout with a consistent accent style (title bar,
rule, bullets, slide numbers) so the deck looks designed without shipping a
binary template.
"""

from io import BytesIO

from contextgit.documents.errors import INSTALL_HINT, MissingRenderer
from contextgit.documents.models import DocumentContent
from contextgit.documents.styles import Style


def render_pptx(document: DocumentContent, style: Style) -> bytes:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - guarded by available()
        raise MissingRenderer(INSTALL_HINT) from exc

    accent = RGBColor.from_string(style.accent.lstrip("#").upper())
    ink = RGBColor.from_string(style.ink.lstrip("#").upper())
    muted = RGBColor.from_string(style.muted.lstrip("#").upper())
    body_font = style.body_font.split(",")[0].strip("'\" ")
    head_font = style.heading_font.split(",")[0].strip("'\" ")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    width = int(presentation.slide_width or 0)
    height = int(presentation.slide_height or 0)

    def rect(slide: object, x: object, y: object, w: object, h: object, color: object) -> None:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)  # type: ignore[attr-defined]
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.shadow.inherit = False

    def text(
        slide: object,
        x: object,
        y: object,
        w: object,
        h: object,
        value: str,
        *,
        size: int,
        color: object,
        bold: bool = False,
        font: str = body_font,
        align: object = None,
    ) -> object:
        box = slide.shapes.add_textbox(x, y, w, h)  # type: ignore[attr-defined]
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font
        if align is not None:
            paragraph.alignment = align
        return box

    # ---- title slide ----
    cover = presentation.slides.add_slide(blank)
    rect(cover, 0, 0, width, Inches(0.28), accent)
    text(cover, Inches(0.9), Inches(2.3), width - Inches(1.8), Inches(1.5), document.title,
         size=40, color=ink, bold=True, font=head_font)
    if document.subtitle:
        text(cover, Inches(0.9), Inches(3.8), width - Inches(1.8), Inches(0.9),
             document.subtitle, size=18, color=muted, font=head_font)
    rect(cover, Inches(0.9), Inches(4.7), Inches(1.2), Inches(0.06), accent)
    text(cover, Inches(0.9), Inches(5.1), width - Inches(1.8), Inches(0.5),
         f"{style.label} · ContextGit", size=12, color=muted, font=head_font)

    total = max(1, len(document.slides))
    for index, slide_data in enumerate(document.slides, start=1):
        page = presentation.slides.add_slide(blank)
        rect(page, 0, 0, width, Inches(0.18), accent)
        text(page, Inches(0.9), Inches(0.7), width - Inches(1.8), Inches(0.9),
             slide_data.title or " ", size=26, color=ink, bold=True, font=head_font)
        rect(page, Inches(0.9), Inches(1.68), Inches(1.0), Inches(0.05), accent)
        box = page.shapes.add_textbox(
            Inches(0.95), Inches(2.05), width - Inches(1.9), height - Inches(2.9)
        )
        frame = box.text_frame
        frame.word_wrap = True
        for bullet_index, bullet in enumerate(slide_data.bullets or [""]):
            paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"•  {bullet}"
            run.font.size = Pt(16)
            run.font.color.rgb = ink
            run.font.name = body_font
            paragraph.space_after = Pt(8)
        text(page, width - Inches(1.5), height - Inches(0.62), Inches(1.0), Inches(0.4),
             f"{index} / {total}", size=10, color=muted, font=head_font, align=PP_ALIGN.RIGHT)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
